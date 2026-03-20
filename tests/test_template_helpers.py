"""Tests for openpyxl template helper functions in api.lib.rfq_pipeline."""

from datetime import date
import pandas as pd
import pytest
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from api.lib.rfq_pipeline import (
    detect_header_row_in_template,
    read_template_headers,
    update_cell_a5,
    write_dataframe,
    add_need_by_cf,
    enable_autofit_wrapped_rows,
)


# ── detect_header_row_in_template ──────────────────────────────

class TestDetectHeaderRowInTemplate:
    def test_finds_header_at_row_7(self, template_workbook):
        ws = template_workbook.active
        known = {"RFQ #", "Customer Company", "Need By Date", "Creation Date",
                 "Sales Account Manager"}
        assert detect_header_row_in_template(ws, known) == 7

    def test_fallback_rfq_token(self):
        """When min_matches is not met, falls back to row containing 'RFQ #'."""
        wb = Workbook()
        ws = wb.active
        ws.cell(row=3, column=1, value="RFQ #")
        ws.cell(row=3, column=2, value="Something Else")
        # Only 1 known header — not enough for min_matches=4
        known = {"RFQ #"}
        result = detect_header_row_in_template(ws, known, min_matches=4)
        assert result == 3

    def test_raises_on_empty_workbook(self):
        wb = Workbook()
        ws = wb.active
        known = {"RFQ #", "Customer Company", "Need By Date", "Creation Date"}
        with pytest.raises(RuntimeError, match="Could not detect header row"):
            detect_header_row_in_template(ws, known)

    def test_multiple_partial_rows_picks_correct(self):
        """Header row with >= min_matches wins even if other rows have some matches."""
        wb = Workbook()
        ws = wb.active
        # Row 2: only 2 matches
        ws.cell(row=2, column=1, value="RFQ #")
        ws.cell(row=2, column=2, value="Customer Company")
        # Row 5: 4 matches (meets threshold)
        ws.cell(row=5, column=1, value="RFQ #")
        ws.cell(row=5, column=2, value="Customer Company")
        ws.cell(row=5, column=3, value="Need By Date")
        ws.cell(row=5, column=4, value="Creation Date")
        known = {"RFQ #", "Customer Company", "Need By Date", "Creation Date"}
        assert detect_header_row_in_template(ws, known, min_matches=4) == 5


# ── read_template_headers ──────────────────────────────────────

class TestReadTemplateHeaders:
    def test_returns_correct_tuples(self, template_workbook):
        ws = template_workbook.active
        headers = read_template_headers(ws, 7)
        # Should be list of (col_idx, header_name) tuples
        assert len(headers) == 17
        assert headers[0] == (1, "RFQ #")
        assert headers[1] == (2, "Customer Company")
        assert headers[2] == (3, "Need By Date")

    def test_strips_trailing_blanks(self):
        wb = Workbook()
        ws = wb.active
        ws.cell(row=1, column=1, value="A")
        ws.cell(row=1, column=2, value="B")
        ws.cell(row=1, column=3, value="")
        ws.cell(row=1, column=4, value="")
        headers = read_template_headers(ws, 1)
        assert len(headers) == 2
        assert headers[-1][1] == "B"


# ── update_cell_a5 ─────────────────────────────────────────────

class TestUpdateCellA5:
    def test_replaces_mmm_dd_yyyy(self):
        wb = Workbook()
        ws = wb.active
        ws["A5"] = "Open RFQ Report Mar 10 2026"
        update_cell_a5(ws, "Mar 23 2026")
        assert ws["A5"].value == "Open RFQ Report Mar 23 2026"

    def test_replaces_slash_date(self):
        wb = Workbook()
        ws = wb.active
        ws["A5"] = "Report 3/10/2026"
        update_cell_a5(ws, "Mar 23 2026")
        assert ws["A5"].value == "Report Mar 23 2026"

    def test_replaces_iso_date(self):
        wb = Workbook()
        ws = wb.active
        ws["A5"] = "Report 2026-03-10"
        update_cell_a5(ws, "Mar 23 2026")
        assert ws["A5"].value == "Report Mar 23 2026"

    def test_appends_when_no_date_pattern(self):
        wb = Workbook()
        ws = wb.active
        ws["A5"] = "Open RFQ Report"
        update_cell_a5(ws, "Mar 23 2026")
        assert "Mar 23 2026" in ws["A5"].value

    def test_non_string_value(self):
        wb = Workbook()
        ws = wb.active
        ws["A5"] = 12345
        update_cell_a5(ws, "Mar 23 2026")
        assert ws["A5"].value == "Mar 23 2026"

    def test_none_value(self):
        wb = Workbook()
        ws = wb.active
        ws["A5"] = None
        update_cell_a5(ws, "Mar 23 2026")
        assert ws["A5"].value == "Mar 23 2026"


# ── write_dataframe ────────────────────────────────────────────

class TestWriteDataframe:
    def test_populates_cells(self, template_workbook):
        ws = template_workbook.active
        template_headers = read_template_headers(ws, 7)
        df = pd.DataFrame({
            "RFQ #": ["1001", "1002"],
            "Customer Company": ["Acme", "Bolt"],
            "Need By Date": ["2026-04-01", "2026-03-15"],
        })
        write_dataframe(ws, df, template_headers, start_row=8)
        assert ws.cell(row=8, column=1).value == "1001"
        assert ws.cell(row=8, column=2).value == "Acme"
        assert ws.cell(row=9, column=1).value == "1002"
        assert ws.cell(row=9, column=2).value == "Bolt"

    def test_preserves_font_from_header(self, template_workbook):
        ws = template_workbook.active
        template_headers = read_template_headers(ws, 7)
        df = pd.DataFrame({"RFQ #": ["1001"]})
        write_dataframe(ws, df, template_headers, start_row=8)
        # Data cell should have the style from the prototype row (row 8 originally)
        cell = ws.cell(row=8, column=1)
        assert cell.font.name == "Calibri"

    def test_handles_empty_dataframe(self, template_workbook):
        ws = template_workbook.active
        template_headers = read_template_headers(ws, 7)
        df = pd.DataFrame(columns=["RFQ #", "Customer Company"])
        write_dataframe(ws, df, template_headers, start_row=8)
        # Should not write anything — row 8 should remain as template placeholder
        assert ws.cell(row=8, column=1).value == ""


# ── add_need_by_cf ─────────────────────────────────────────────

class TestAddNeedByCf:
    def test_creates_two_rules(self, template_workbook):
        ws = template_workbook.active
        template_headers = read_template_headers(ws, 7)
        add_need_by_cf(ws, template_headers, data_start=8, n_rows=5,
                       monday=date(2026, 3, 23))
        # Two rules (gold + orange) are added; they may share a CF range entry
        total_rules = sum(len(cf.rules) for cf in ws.conditional_formatting)
        assert total_rules == 2

    def test_cf_range_excludes_column_a(self, template_workbook):
        ws = template_workbook.active
        template_headers = read_template_headers(ws, 7)
        add_need_by_cf(ws, template_headers, data_start=8, n_rows=5,
                       monday=date(2026, 3, 23))
        for cf in ws.conditional_formatting:
            for cell_range in cf.cells.ranges:
                assert str(cell_range).startswith("B"), \
                    f"CF range should start at B, got {cell_range}"

    def test_skips_on_zero_rows(self, template_workbook):
        ws = template_workbook.active
        template_headers = read_template_headers(ws, 7)
        initial_count = len(ws.conditional_formatting._cf_rules)
        add_need_by_cf(ws, template_headers, data_start=8, n_rows=0,
                       monday=date(2026, 3, 23))
        assert len(ws.conditional_formatting._cf_rules) == initial_count

    def test_skips_when_no_need_by_header(self):
        wb = Workbook()
        ws = wb.active
        ws.cell(row=1, column=1, value="RFQ #")
        headers = [(1, "RFQ #")]
        initial_count = len(ws.conditional_formatting._cf_rules)
        add_need_by_cf(ws, headers, data_start=2, n_rows=3,
                       monday=date(2026, 3, 23), need_by_header="Need By Date")
        assert len(ws.conditional_formatting._cf_rules) == initial_count


# ── enable_autofit_wrapped_rows ────────────────────────────────

class TestEnableAutofitWrappedRows:
    def test_sets_wrap_text(self):
        wb = Workbook()
        ws = wb.active
        for r in range(2, 5):
            for c in range(1, 18):
                ws.cell(row=r, column=c, value=f"R{r}C{c}")
        enable_autofit_wrapped_rows(ws, start_row=2, n_rows=3)
        for r in range(2, 5):
            for c in range(1, 18):  # A=1 to Q=17
                assert ws.cell(row=r, column=c).alignment.wrap_text is True

    def test_resets_row_height_to_none(self):
        wb = Workbook()
        ws = wb.active
        for r in range(2, 5):
            ws.row_dimensions[r].height = 30
            for c in range(1, 18):
                ws.cell(row=r, column=c, value="x")
        enable_autofit_wrapped_rows(ws, start_row=2, n_rows=3)
        for r in range(2, 5):
            assert ws.row_dimensions[r].height is None

    def test_skips_on_zero_rows(self):
        wb = Workbook()
        ws = wb.active
        ws.cell(row=2, column=1, value="x")
        ws.row_dimensions[2].height = 30
        enable_autofit_wrapped_rows(ws, start_row=2, n_rows=0)
        assert ws.row_dimensions[2].height == 30
