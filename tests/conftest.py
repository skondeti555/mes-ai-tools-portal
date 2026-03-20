import pytest
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment


@pytest.fixture
def sample_master_df():
    """A minimal DataFrame matching FINAL_COLUMNS schema with 6 rows across 4 countries."""
    return pd.DataFrame({
        "RFQ #": ["1001", "1002", "1003", "1004", "1005", "1006"],
        "Customer Company": ["Acme", "Bolt", "Crane", "Delta", "Echo", "Foxtrot"],
        "Need By Date": ["2026-04-01", "2026-03-15", "", "2026-05-01", "2026-03-20", "2026-04-10"],
        "Creation Date": ["2026-01-01"] * 6,
        "Sales Account Manager": ["Smith"] * 6,
        "RFQ Coordinator": ["Jones"] * 6,
        "Project Name": ["P1", "P2", "P3", "P4", "P5", "P6"],
        "Invited Supplier Count": ["3", "2", "5", "1", "4", "2"],
        "Quoted Supplier Count": ["1", "0", "2", "0", "1", "1"],
        "Commodity": ["Steel"] * 6,
        "Process": ["CNC"] * 6,
        "RFQ Type / Classification": ["Standard"] * 6,
        "Total Parts": ["10", "20", "5", "15", "8", "12"],
        "Total Quantity": ["100", "200", "50", "150", "80", "120"],
        "Floated Country": ["IN", "CN", "MX", "VN", "IN, CN", "MX;VN"],
        "Open Floated Country": ["IN", "CN", "MX", "VN", "IN, CN", "MX;VN"],
        "Standards Requirement": ["ISO"] * 6,
        "Internal Comments": [""] * 6,
    })


@pytest.fixture
def sample_master_df_with_supplier(sample_master_df):
    """Master DF with _InvitedSupplier column for Mexico exclusion testing."""
    df = sample_master_df.copy()
    df["_InvitedSupplier"] = ["Supplier A", "Supplier B", "Metrics Works Saltillo",
                               "Supplier D", "Supplier E", "Supplier F"]
    return df


@pytest.fixture
def template_workbook():
    """A minimal openpyxl Workbook mimicking a country template."""
    wb = Workbook()
    ws = wb.active
    # Row 5 = date cell
    ws["A5"] = "Open RFQ Report Mar 10 2026"
    # Row 7 = header row
    headers = [
        "RFQ #", "Customer Company", "Need By Date", "Creation Date",
        "Sales Account Manager", "RFQ Coordinator", "Project Name",
        "Invited Supplier Count", "Quoted Supplier Count", "Commodity",
        "Process", "RFQ Type / Classification", "Total Parts",
        "Total Quantity", "Floated Country", "Open Floated Country",
        "Standards Requirement",
    ]
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=7, column=col_idx, value=h)
        cell.font = Font(bold=True, name="Calibri", size=11)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    # Row 8 = blank styled data row
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=8, column=col_idx, value="")
        cell.font = Font(name="Calibri", size=11)
        cell.alignment = Alignment(wrap_text=True)
    return wb
